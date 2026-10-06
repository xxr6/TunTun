"""文档解析：按扩展名抽取纯文本。

支持 TXT / Markdown / PDF / DOCX / PPTX / EPUB。
DOCX、PPTX、EPUB 用延迟导入：不解析这些格式时不强依赖。

技术文档 7.3 的取舍：不引 markitdown（依赖重），每个格式用最精准的小库——
python-docx（段落+表格）、python-pptx（逐页文本框）、ebooklib（EPUB 内 HTML 剥标签）。
"""
import html as _html
import io
import re

# 支持直接解析的扩展名
SUPPORTED = {"txt", "md", "markdown", "pdf", "docx", "pptx", "epub"}


def extract_text(filename: str, data: bytes) -> tuple[str, int | None]:
    """返回 (纯文本, 页数)。页数对 PDF/PPTX 有效，其余为 None。"""
    name = filename or ""
    ext = (name.rsplit(".", 1)[-1] if "." in name else "").lower()

    if ext in ("txt", "md", "markdown"):
        return data.decode("utf-8", errors="replace"), None

    if ext == "pdf":
        import pymupdf  # 延迟导入：非 PDF 场景不强依赖

        doc = pymupdf.open(stream=data, filetype="pdf")
        try:
            pages = [page.get_text() for page in doc]
            return "\n\n".join(pages), len(pages)
        finally:
            doc.close()

    if ext == "docx":
        return _extract_docx(data), None

    if ext == "pptx":
        return _extract_pptx(data)

    if ext == "epub":
        return _extract_epub(data), None

    raise ValueError(f"暂不支持解析 .{ext} 格式，已支持：{', '.join(sorted(SUPPORTED))}")


def _extract_docx(data: bytes) -> str:
    from docx import Document as DocxDocument

    d = DocxDocument(io.BytesIO(data))
    parts = [p.text.strip() for p in d.paragraphs if p.text.strip()]
    # 表格内容也抽出来（需求文档里的对比表很常见）
    for table in d.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n\n".join(parts)


def _extract_pptx(data: bytes) -> tuple[str, int | None]:
    from pptx import Presentation

    prs = Presentation(io.BytesIO(data))
    slides: list[str] = []
    for i, slide in enumerate(prs.slides, 1):
        texts = []
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False):
                t = shape.text_frame.text.strip()
                if t:
                    texts.append(t)
        if texts:
            slides.append(f"【第 {i} 页】\n" + "\n".join(texts))
    return "\n\n".join(slides), len(prs.slides)


_TAG_BLOCK = re.compile(r"<(script|style)[\s\S]*?</\1>", re.I)
_TAG = re.compile(r"<[^>]+>")


def _extract_epub(data: bytes) -> str:
    from ebooklib import ITEM_DOCUMENT, epub

    book = epub.read_epub(io.BytesIO(data), options={"ignore_ncx": True})
    parts: list[str] = []
    for item in book.get_items_of_type(ITEM_DOCUMENT):
        raw = item.get_content().decode("utf-8", errors="replace")
        text = _TAG_BLOCK.sub(" ", raw)
        text = _TAG.sub("\n", text)  # 块级标签直接断行
        text = _html.unescape(text)
        lines = [ln.strip() for ln in text.splitlines()]
        block = "\n".join(ln for ln in lines if ln)
        if block.strip():
            parts.append(block)
    return "\n\n".join(parts)
